"""Importerer et eksisterende kvartalsregneark ind i emballage-modulet.

Bruges til at få historikken med, før modulet tages i brug. Arket findes i to
layouts – det ældste faneblad har kolonnerne "Flaske type / Antal flasker /
Leveringsdato …", de nyere har "Dato / Faktura / Kvartal / Vare / Antal …".
Begge understøttes.

Linjerne grupperes til følgesedler pr. (fakturanummer, dato). Vægtene tages fra
arket, ikke fra stamdata: arket er kilden, og så kan man bagefter se om de to
er uenige.

    venv/bin/python -m scripts.import_kvartalsark ark.xlsx --dry-run
    venv/bin/python -m scripts.import_kvartalsark ark.xlsx
"""

import argparse
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl                                              # noqa: E402

from config.config import sessionLocal                       # noqa: E402
from config.model import (DEFAULT_CARTON_KG, PackagingInvoice,  # noqa: E402
                          PackagingLine, PackagingMaterial)
from services.packaging_excel import quarter_of              # noqa: E402
from services.packaging_pdf import match_material            # noqa: E402

# Det ældste faneblad har ingen Faktura-kolonne. Nummeret står på den
# tilhørende pakirna lista ("ROT-EUR 41/01/011 od 23.09.2025."), og de 22
# linjer stemmer med arket, så det udfyldes her.
KENDTE_FAKTURANUMRE = {"Q3 2025": "41/01/011"}

GAMMELT_LAYOUT = {
    "name": "Flaske type", "qty": "Antal flasker", "date": "Leveringsdato",
    "glass": "kg pr. flaske (glas)", "carton": "kg pap pr. flaske",
    "quarter": "Kvartal", "invoice": None, "note": None,
}
NYT_LAYOUT = {
    "name": "Vare", "qty": "Antal", "date": "Dato", "glass": "Tom flaske kg",
    "carton": "Pap pr. flaske kg", "quarter": "Kvartal", "invoice": "Faktura",
    "note": "Note",
}


def parse_dato(v):
    if isinstance(v, datetime):
        return v
    if not v:
        return None
    for fmt in ("%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(v).strip(), fmt)
        except ValueError:
            continue
    return None


def læs_ark(ws) -> list[dict]:
    hoved = [c.value for c in ws[1]]
    layout = NYT_LAYOUT if "Vare" in hoved else GAMMELT_LAYOUT
    idx = {k: (hoved.index(v) + 1 if v and v in hoved else None)
           for k, v in layout.items()}

    linjer = []
    for r in range(2, ws.max_row + 1):
        navn = ws.cell(r, idx["name"]).value
        if not navn or str(navn).strip().upper() == "TOTAL":
            continue
        antal = ws.cell(r, idx["qty"]).value
        if not isinstance(antal, (int, float)):
            continue
        glas = ws.cell(r, idx["glass"]).value
        pap = ws.cell(r, idx["carton"]).value
        linjer.append({
            "item_name": str(navn).strip(),
            "quantity": int(antal),
            "glass_kg": float(glas) if isinstance(glas, (int, float)) else None,
            "carton_kg": float(pap) if isinstance(pap, (int, float))
            else DEFAULT_CARTON_KG,
            "invoice_date": parse_dato(ws.cell(r, idx["date"]).value),
            "invoice_number": (ws.cell(r, idx["invoice"]).value
                               if idx["invoice"] else None),
            "quarter": (ws.cell(r, idx["quarter"]).value
                        if idx["quarter"] else ws.title),
            "note": (ws.cell(r, idx["note"]).value if idx["note"] else None),
        })
    return linjer


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("fil")
    ap.add_argument("--dry-run", action="store_true",
                    help="vis hvad der ville blive oprettet, uden at gemme")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.fil, data_only=True)
    db = sessionLocal()
    try:
        materialer = db.query(PackagingMaterial).all()
        if not materialer:
            print("Der er ingen emballage-stamdata. Kør først:")
            print("  venv/bin/python -m scripts.seed_packaging_materials")
            return 1

        # Grupper linjer til følgesedler pr. (fakturanummer, dato)
        følgesedler: dict = {}
        advarsler = []
        for ws in wb.worksheets:
            if ws.title.lower().startswith("oversigt"):
                continue
            for linje in læs_ark(ws):
                nummer = (linje["invoice_number"]
                          or KENDTE_FAKTURANUMRE.get(linje["quarter"]))
                nøgle = (nummer, linje["invoice_date"])
                fs = følgesedler.setdefault(nøgle, {
                    "invoice_number": nummer,
                    "invoice_date": linje["invoice_date"],
                    "quarter": linje["quarter"] or quarter_of(linje["invoice_date"]),
                    "lines": [],
                })

                mat, score, hvordan = match_material(
                    {"article_code": None, "item_name": linje["item_name"]},
                    materialer)
                if mat is None:
                    advarsler.append(
                        f"{linje['item_name']}: ingen stamdata (bedste match {score})")
                elif (linje["glass_kg"] is not None and mat.glass_kg is not None
                      and abs(float(mat.glass_kg) - linje["glass_kg"]) > 1e-6):
                    advarsler.append(
                        f"{linje['item_name']}: arket siger {linje['glass_kg']} kg, "
                        f"stamdata siger {float(mat.glass_kg)} kg – arket bruges")

                fs["lines"].append({
                    **linje,
                    "material_id": mat.id if mat else None,
                    "article_code": mat.article_code if mat else None,
                    "matched": mat.name if mat else None,
                })

        # ---- vis planen ----
        print(f"{len(følgesedler)} følgesedler fra {args.fil}\n")
        i_alt_linjer = i_alt_flasker = 0
        for (nummer, dato), fs in sorted(
                følgesedler.items(), key=lambda kv: (kv[0][1] or datetime.min)):
            flasker = sum(l["quantity"] for l in fs["lines"])
            glas = sum((l["glass_kg"] or 0) * l["quantity"] for l in fs["lines"])
            umatchede = sum(1 for l in fs["lines"] if not l["material_id"])
            print(f"  {fs['quarter']:<9} {nummer or '(uden nummer)':<12} "
                  f"{dato:%d-%m-%Y}  {len(fs['lines']):>2} linjer  "
                  f"{flasker:>4} flasker  {glas:>8.3f} kg glas"
                  + (f"  ({umatchede} uden stamdata)" if umatchede else ""))
            i_alt_linjer += len(fs["lines"])
            i_alt_flasker += flasker
        print(f"\n  I alt: {i_alt_linjer} linjer, {i_alt_flasker} flasker")

        if advarsler:
            print(f"\n  {len(advarsler)} bemærkninger:")
            for a in sorted(set(advarsler)):
                print("   -", a)

        # ---- eksisterende følgesedler springes over ----
        findes = {i.invoice_number for i in db.query(PackagingInvoice).all()
                  if i.invoice_number}
        springes_over = [n for (n, _) in følgesedler if n in findes]
        if springes_over:
            print(f"\n  Findes allerede, springes over: {springes_over}")

        if args.dry_run:
            print("\n(--dry-run: der er ikke gemt noget)")
            return 0

        oprettet = 0
        for (nummer, dato), fs in følgesedler.items():
            if nummer and nummer in findes:
                continue
            faktura = PackagingInvoice(
                invoice_number=nummer,
                invoice_date=fs["invoice_date"],
                quarter=fs["quarter"],
                supplier="Galic",
                notes="Importeret fra kvartalsregnearket",
            )
            for n, l in enumerate(fs["lines"], start=1):
                faktura.lines.append(PackagingLine(
                    line_no=n,
                    article_code=l["article_code"],
                    item_name=l["item_name"],
                    quantity=l["quantity"],
                    glass_kg=l["glass_kg"],
                    carton_kg=l["carton_kg"],
                    material_id=l["material_id"],
                    note=l["note"],
                ))
            db.add(faktura)
            oprettet += 1
        db.commit()
        print(f"\n{oprettet} følgesedler oprettet.")
        return 0
    finally:
        db.close()


if __name__ == "__main__":
    sys.exit(main())
